"""
Google Docs DOCX -> Obsidian Markdown Converter
------------------------------------------------
Converts .docx files (exported from Google Docs) into Markdown files
compatible with Obsidian.

Handles two export formats automatically:

  HEADING MODE: Documents where Google Docs tabs exported as proper
  Word heading styles (Heading 1, Heading 2, etc.). Produces nested
  folders and files mirroring the heading hierarchy, with [[wikilinks]]
  from parent notes to their direct children.

  SECTION BREAK MODE: Documents where Google Docs tabs exported as
  nextPage section breaks with all content in 'normal' style (the more
  common Google Docs export behaviour). Each section break = one tab =
  one Markdown file. All files are flat inside a folder named after the
  document. The first non-empty paragraph of each section becomes the
  file title.

Rich formatting is preserved in both modes:
  - Bold -> **text**, Italic -> *text*, Bold+Italic -> ***text***
  - Bullet lists and numbered lists with nesting
  - Hyperlinks -> [text](url)

SINGLE FOLDER MODE:
    python docx_to_markdown.py <input_folder> <output_folder>

BATCH MODE (subfolders):
    python docx_to_markdown.py <input_folder> <output_folder> --batch

Requirements: Python 3.6+, python-docx (auto-installed on first run).
"""

import os
import sys
import re
import subprocess


# ---------------------------------------------------------------------------
# Auto-install python-docx
# ---------------------------------------------------------------------------

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    print("Installing python-docx (required once)...")
    subprocess.check_call(
        [sys.executable, '-m', 'pip', 'install', 'python-docx',
         '--break-system-packages'],
        stdout=subprocess.DEVNULL
    )
    from docx import Document
    from docx.oxml.ns import qn


# ---------------------------------------------------------------------------
# Paragraph inspection helpers
# ---------------------------------------------------------------------------

def get_heading_level(para):
    """Return heading level 1-6, or None if not a heading style."""
    style = para.style.name if para.style else ''
    m = re.match(r'[Hh]eading\s+(\d)', style)
    return int(m.group(1)) if m else None


def is_nextpage_break(para):
    """Return True if paragraph ends with a nextPage section break."""
    pPr = para._p.find(qn('w:pPr'))
    if pPr is None:
        return False
    sectPr = pPr.find(qn('w:sectPr'))
    if sectPr is None:
        return False
    sectType = sectPr.find(qn('w:type'))
    break_type = sectType.get(qn('w:val')) if sectType is not None else ''
    return break_type == 'nextPage'


def get_list_info(para):
    """
    Return (indent_level, is_bullet) for list paragraphs, or None.
    Handles both style-name-based (Google Docs) and numPr-based (Word) lists.
    """
    style_name = para.style.name if para.style else ''

    bullet_match = re.match(r'List Bullet\s*(\d*)', style_name)
    number_match = re.match(r'List Number\s*(\d*)', style_name)

    if bullet_match:
        level = int(bullet_match.group(1)) - 1 if bullet_match.group(1) else 0
        return (max(0, level), True)
    if number_match:
        level = int(number_match.group(1)) - 1 if number_match.group(1) else 0
        return (max(0, level), False)

    pPr = para._p.find(qn('w:pPr'))
    if pPr is not None:
        numPr = pPr.find(qn('w:numPr'))
        if numPr is not None:
            ilvl_el = numPr.find(qn('w:ilvl'))
            level = int(ilvl_el.get(qn('w:val'))) if ilvl_el is not None else 0
            is_bullet = 'Number' not in style_name
            return (level, is_bullet)

    return None


def runs_to_markdown(para):
    """Convert paragraph runs to markdown with bold, italic, and hyperlinks."""
    result = ''
    for child in para._p:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag == 'r':
            rPr   = child.find(qn('w:rPr'))
            bold  = rPr is not None and rPr.find(qn('w:b'))  is not None
            italic= rPr is not None and rPr.find(qn('w:i'))  is not None
            t_el  = child.find(qn('w:t'))
            text  = (t_el.text or '') if t_el is not None else ''
            if not text:
                continue
            if bold and italic:
                text = f'***{text}***'
            elif bold:
                text = f'**{text}**'
            elif italic:
                text = f'*{text}*'
            result += text

        elif tag == 'hyperlink':
            r_id = child.get(qn('r:id'))
            url  = ''
            if r_id and hasattr(para.part, 'rels') and r_id in para.part.rels:
                url = para.part.rels[r_id].target_ref
            link_text = ''
            for run_el in child.findall(qn('w:r')):
                t_el  = run_el.find(qn('w:t'))
                if t_el is None or not t_el.text:
                    continue
                rPr   = run_el.find(qn('w:rPr'))
                bold  = rPr is not None and rPr.find(qn('w:b'))  is not None
                italic= rPr is not None and rPr.find(qn('w:i'))  is not None
                text  = t_el.text
                if bold and italic:
                    text = f'***{text}***'
                elif bold:
                    text = f'**{text}**'
                elif italic:
                    text = f'*{text}*'
                link_text += text
            if url and link_text:
                result += f'[{link_text}]({url})'
            elif link_text:
                result += link_text

    return result


def para_to_markdown(para, skip_heading_prefix=False):
    """Convert a single paragraph to a Markdown string."""
    heading_level = get_heading_level(para)
    if heading_level is not None and not skip_heading_prefix:
        text = para.text.strip()
        return f"{'#' * heading_level} {text}" if text else ''

    list_info = get_list_info(para)
    text = runs_to_markdown(para)

    if not text.strip():
        return ''

    if list_info is not None:
        level, is_bullet = list_info
        indent  = '  ' * level
        prefix  = '-' if is_bullet else '1.'
        return f'{indent}{prefix} {text}'

    return text


def paragraphs_to_markdown_body(paragraphs):
    """Convert a list of paragraphs to a markdown body string."""
    lines = []
    prev_was_list = False

    for para in paragraphs:
        md = para_to_markdown(para)
        is_list = get_list_info(para) is not None

        if md and lines:
            if not (is_list and prev_was_list):
                lines.append('')

        if md:
            lines.append(md)
        prev_was_list = is_list

    return '\n'.join(lines).strip()


# ---------------------------------------------------------------------------
# Shared utilities
# ---------------------------------------------------------------------------

def yaml_safe_title(title):
    escaped = title.replace('\\', '\\\\').replace('"', '\\"')
    return f'"{escaped}"'


def build_front_matter(title):
    return f'---\ntitle: {yaml_safe_title(title)}\n---\n\n'


def safe_name(name):
    """Strip characters invalid in Windows file/folder names."""
    name = re.sub(r'[\\/*?:"<>|]', '_', name)
    return name[:80].strip()


def title_from_content(paragraphs):
    """
    Derive a title from the first non-empty paragraph.
    Uses full text if it looks like a short title (<=60 chars, no period mid-text),
    otherwise falls back to first 3 words or 15 chars + '...'.
    """
    for para in paragraphs:
        text = para.text.strip()
        if text:
            if len(text) <= 60 and not re.search(r'\.\s+\w', text):
                return text
            words = text.split()
            by_words = ' '.join(words[:3])
            by_chars = text[:15]
            excerpt = by_words if len(by_words) <= len(by_chars) else by_chars
            return excerpt.strip() + '...'
    return 'Untitled'


# ---------------------------------------------------------------------------
# HEADING MODE — documents with proper Heading styles
# ---------------------------------------------------------------------------

class Section:
    def __init__(self, level, title, parent=None):
        self.level      = level
        self.title      = title
        self.paragraphs = []
        self.children   = []
        self.parent     = parent


def build_section_tree(doc, root_title):
    root = Section(level=0, title=root_title)
    current_by_level = {0: root}

    for para in doc.paragraphs:
        heading_level = get_heading_level(para)
        if heading_level is not None:
            title = para.text.strip()
            if not title:
                continue
            parent_level = heading_level - 1
            while parent_level > 0 and parent_level not in current_by_level:
                parent_level -= 1
            parent = current_by_level.get(parent_level, root)
            section = Section(level=heading_level, title=title, parent=parent)
            parent.children.append(section)
            current_by_level[heading_level] = section
            for lvl in list(current_by_level.keys()):
                if lvl > heading_level:
                    del current_by_level[lvl]
        else:
            deepest = max(current_by_level.values(), key=lambda s: s.level)
            deepest.paragraphs.append(para)

    return root


def section_to_markdown(section):
    lines = []
    prev_was_list = False

    for para in section.paragraphs:
        md = para_to_markdown(para)
        is_list = get_list_info(para) is not None
        if md and lines:
            if not (is_list and prev_was_list):
                lines.append('')
        if md:
            lines.append(md)
        prev_was_list = is_list

    if section.children:
        if lines:
            lines.append('')
        lines.append('## Contents')
        for child in section.children:
            lines.append(f'- [[{child.title}]]')

    return '\n'.join(lines).strip()


def write_heading_tree(section, output_path):
    os.makedirs(output_path, exist_ok=True)
    for child in section.children:
        child_name = safe_name(child.title)
        content = build_front_matter(child.title) + section_to_markdown(child)
        if child.children:
            child_folder = os.path.join(output_path, child_name)
            os.makedirs(child_folder, exist_ok=True)
            with open(os.path.join(child_folder, f'{child_name}.md'),
                      'w', encoding='utf-8') as f:
                f.write(content)
            write_heading_tree(child, child_folder)
        else:
            with open(os.path.join(output_path, f'{child_name}.md'),
                      'w', encoding='utf-8') as f:
                f.write(content)


def convert_heading_mode(doc, doc_name, doc_output):
    root = build_section_tree(doc, doc_name)
    if root.paragraphs:
        pre_content = build_front_matter(doc_name) + section_to_markdown(root)
        with open(os.path.join(doc_output, f'{safe_name(doc_name)}.md'),
                  'w', encoding='utf-8') as f:
            f.write(pre_content)
    write_heading_tree(root, doc_output)
    return sum(1 for _ in root.children)


# ---------------------------------------------------------------------------
# SECTION BREAK MODE — documents with nextPage breaks (Google Docs tabs)
# ---------------------------------------------------------------------------

def split_by_section_breaks(doc):
    """
    Split document paragraphs into sections at nextPage breaks.
    Returns list of paragraph lists, one per section/tab.
    """
    sections = []
    current = []
    for para in doc.paragraphs:
        current.append(para)
        if is_nextpage_break(para):
            sections.append(current)
            current = []
    if current:
        sections.append(current)
    return sections


def convert_section_break_mode(doc, doc_name, doc_output):
    sections = split_by_section_breaks(doc)
    file_count = 0
    used_names = {}

    for section_paras in sections:
        # Skip entirely empty sections
        if not any(p.text.strip() for p in section_paras):
            continue

        title = title_from_content(section_paras)
        body  = paragraphs_to_markdown_body(section_paras)

        # Skip sections that are just separators or truly empty after conversion
        if not body.strip():
            continue

        content   = build_front_matter(title) + body
        base_name = safe_name(title)

        # Handle filename collisions
        if base_name in used_names:
            used_names[base_name] += 1
            file_name = f'{base_name}_{used_names[base_name]}.md'
        else:
            used_names[base_name] = 0
            file_name = f'{base_name}.md'

        with open(os.path.join(doc_output, file_name), 'w', encoding='utf-8') as f:
            f.write(content)
        file_count += 1

    return file_count


# ---------------------------------------------------------------------------
# Document mode detection and per-file conversion
# ---------------------------------------------------------------------------

def has_heading_styles(doc):
    """Return True if the document uses Word heading styles."""
    for para in doc.paragraphs:
        if get_heading_level(para) is not None:
            return True
    return False


def has_section_breaks(doc):
    """Return True if the document uses nextPage section breaks."""
    for para in doc.paragraphs:
        if is_nextpage_break(para):
            return True
    return False


def convert_docx_file(input_path, output_folder):
    """
    Convert one .docx file. Auto-detects heading vs section-break mode.
    """
    doc_name   = os.path.splitext(os.path.basename(input_path))[0]
    doc_output = os.path.join(output_folder, safe_name(doc_name))
    os.makedirs(doc_output, exist_ok=True)

    doc = Document(input_path)

    if has_heading_styles(doc):
        mode  = 'heading'
        count = convert_heading_mode(doc, doc_name, doc_output)
    elif has_section_breaks(doc):
        mode  = 'section-break'
        count = convert_section_break_mode(doc, doc_name, doc_output)
    else:
        # Plain document with no structure — write as single file
        mode    = 'plain'
        body    = paragraphs_to_markdown_body(doc.paragraphs)
        content = build_front_matter(doc_name) + body
        with open(os.path.join(doc_output, f'{safe_name(doc_name)}.md'),
                  'w', encoding='utf-8') as f:
            f.write(content)
        count = 1

    return mode, count


# ---------------------------------------------------------------------------
# Folder / batch handling
# ---------------------------------------------------------------------------

def process_folder(input_folder, output_folder):
    docx_files = sorted(f for f in os.listdir(input_folder)
                        if f.lower().endswith('.docx'))
    if not docx_files:
        print("  No .docx files found.")
        return 0, 0

    success = errors = 0
    for i, fname in enumerate(docx_files, 1):
        try:
            mode, count = convert_docx_file(
                os.path.join(input_folder, fname), output_folder)
            print(f"  {i}/{len(docx_files)}: {fname}  [{mode}, {count} files]")
            success += 1
        except Exception as e:
            print(f"  ERROR {fname}: {e}")
            errors += 1
    return success, errors


def convert_single(input_folder, output_folder):
    docx_files = [f for f in os.listdir(input_folder)
                  if f.lower().endswith('.docx')]
    print(f"Found {len(docx_files)} .docx files. Converting...")
    success, errors = process_folder(input_folder, output_folder)
    print_summary(success, errors, output_folder)


def convert_batch_mode(input_folder, output_folder):
    subfolders = sorted(d for d in os.listdir(input_folder)
                        if os.path.isdir(os.path.join(input_folder, d)))
    root_files = [f for f in os.listdir(input_folder)
                  if f.lower().endswith('.docx')]

    if not subfolders and not root_files:
        print("No subfolders or .docx files found.")
        return

    total_success = total_errors = 0

    if root_files:
        print(f"\nRoot: {len(root_files)} files")
        s, e = process_folder(input_folder, os.path.join(output_folder, '_root'))
        total_success += s
        total_errors  += e

    for subfolder in subfolders:
        sub_input = os.path.join(input_folder, subfolder)
        docx_files = [f for f in os.listdir(sub_input)
                      if f.lower().endswith('.docx')]
        if not docx_files:
            print(f"\nSkipping '{subfolder}' — no .docx files.")
            continue
        print(f"\nBatch '{subfolder}': {len(docx_files)} files")
        s, e = process_folder(sub_input, os.path.join(output_folder, subfolder))
        total_success += s
        total_errors  += e

    print_summary(total_success, total_errors, output_folder)


def print_summary(success, errors, output_folder):
    print(f"\n{'='*60}")
    print(f"CONVERSION COMPLETE")
    print(f"  Converted: {success}")
    if errors:
        print(f"  Errors:    {errors}")
    print(f"  Output:    {output_folder}")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == '__main__':
    batch_mode = '--batch' in sys.argv
    args = [a for a in sys.argv[1:] if not a.startswith('--')]

    if len(args) != 2:
        print(__doc__)
        sys.exit(1)

    input_folder, output_folder = args

    if not os.path.isdir(input_folder):
        print(f"Error: input folder not found: {input_folder}")
        sys.exit(1)

    if batch_mode:
        convert_batch_mode(input_folder, output_folder)
    else:
        convert_single(input_folder, output_folder)
